"""Lấy mẫu 4 x 10 records để kiểm chứng với chuyên gia (Human-in-the-loop, mục 3.5).

- Lấy mẫu phân tầng (stratified) theo nhãn trạng thái / hiệu ứng để phủ đủ các loại.
- Seed cố định -> tái lập được.
- Xuất: samples.json (có index gốc để truy vết) + expert_review_sheet.xlsx (phiếu chấm).

Chạy: python -X utf8 expert_review/sample_for_experts.py [--seed 42] [--n 10]
"""
import argparse
import json
import random
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent

SRC = {
    "finetune": ROOT / "data-finetune" / "finetune.jsonl",
    "generative": ROOT / "data-generative" / "generative.jsonl",
    "nutrition": ROOT / "Data-RAG" / "400 data dinh duong_v3.json",
    "psychology": ROOT / "Data-RAG" / "psychology_dataset_vi.jsonl",
}


def load_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def stratified(records, key_fn, n, rng):
    """Chia n mẫu cho các tầng: mỗi tầng ít nhất 1 (nếu n đủ), phần còn lại theo tỉ lệ."""
    groups = defaultdict(list)
    for idx, r in enumerate(records):
        groups[key_fn(r)].append(idx)
    keys = sorted(groups, key=lambda k: -len(groups[k]))
    alloc = {k: 0 for k in keys}
    for k in keys[:n]:
        alloc[k] = 1
    remaining = n - sum(alloc.values())
    total = sum(len(groups[k]) for k in keys)
    while remaining > 0:
        # tầng đang "thiếu" nhiều nhất so với tỉ lệ thực
        k = max(keys, key=lambda k: len(groups[k]) / total * n - alloc[k]
                if alloc[k] < len(groups[k]) else -1e9)
        alloc[k] += 1
        remaining -= 1
    picked = []
    for k in keys:
        picked += rng.sample(groups[k], alloc[k])
    return sorted(picked)


# ---------- chọn mẫu ----------

def sample_finetune(rng, n):
    recs = load_jsonl(SRC["finetune"])
    # tầng theo độ dài câu trả lời (tứ phân vị) để phủ cả câu ngắn lẫn dài
    lens = sorted(len(r["messages"][2]["content"]) for r in recs)
    q1, q2, q3 = (lens[len(lens) * k // 4] for k in (1, 2, 3))
    def key(r):
        L = len(r["messages"][2]["content"])
        return "Q1_ngan" if L < q1 else "Q2" if L < q2 else "Q3" if L < q3 else "Q4_dai"
    rows = []
    for k, i in enumerate(stratified(recs, key, n, rng), 1):
        r = recs[i]
        rows.append({
            "review_id": f"FT-{k:02d}",
            "source_index": i,  # dòng (0-based) trong finetune.jsonl
            "type": key(r),
            "user": r["messages"][1]["content"],
            "assistant": r["messages"][2]["content"],
        })
    return rows, recs[0]["messages"][0]["content"]


def sample_generative(rng, n):
    recs = load_jsonl(SRC["generative"])
    # chấp nhận cả "MentalState:" (lỗi gõ) và dấu chấm cuối ("Stressed.")
    ms = lambda r: re.search(r"Mental\s*State:\s*([^|]+)", r["_m"]["system"]).group(1).strip().rstrip(".")
    # đọc theo role (11 record có system message bị lặp 2 lần -> 4 message)
    for r in recs:
        by_role = {}
        for m in r["messages"]:
            by_role.setdefault(m["role"], m["content"])
        r["_m"] = by_role
    # khử trùng lặp (file có nhiều dòng giống hệt nhau) -> giữ lần xuất hiện đầu
    seen, ok = set(), []
    for i, r in enumerate(recs):
        user, asst = r["_m"]["user"], r["_m"]["assistant"]
        if (user, asst) in seen:
            continue
        seen.add((user, asst))
        ok.append(i)
    sub = [recs[i] for i in ok]
    rows = []
    for k, j in enumerate(stratified(sub, ms, n, rng), 1):
        i, r = ok[j], sub[j]
        rows.append({
            "review_id": f"GEN-{k:02d}",
            "source_index": i,
            "test_id": r.get("test_id", f"line_{i + 1}"),
            "mental_state": ms(r),
            "context": r["_m"]["system"],
            "user": r["_m"]["user"],
            "assistant": r["_m"]["assistant"],
        })
    return rows


def sample_nutrition(rng, n):
    with open(SRC["nutrition"], encoding="utf-8") as f:
        recs = json.load(f)
    # bỏ record dưới ngưỡng 20% DV (content_text = null, không có gì để kiểm chứng)
    ok = [i for i, r in enumerate(recs) if r["content_text"]]
    sub = [recs[i] for i in ok]
    rows = []
    for k, j in enumerate(stratified(sub, lambda r: r["ontology_effect"], n, rng), 1):
        i, r = ok[j], sub[j]
        rows.append({
            "review_id": f"NUT-{k:02d}",
            "source_index": i,
            "food_item": r["food_item"],
            "key_nutrients": ", ".join(r["key_nutrients"] or []),
            "ontology_effect": r["ontology_effect"],
            "content_text": r["content_text"],
            "source_reference": r["source_reference"],
        })
    return rows


def sample_psychology(rng, n):
    recs = load_jsonl(SRC["psychology"])
    # bỏ chunk quá ngắn (tiêu đề, mảnh vụn) để đánh giá có ý nghĩa
    ok = [i for i, r in enumerate(recs) if len(r["content_text"]) >= 150]
    sub = [recs[i] for i in ok]
    rows = []
    for k, j in enumerate(stratified(sub, lambda r: r["target_mental_state"], n, rng), 1):
        r = sub[j]
        rows.append({
            "review_id": f"PSY-{k:02d}",
            "source_index": ok[j],
            "chunk_id": r["chunk_id"],
            "target_mental_state": r["target_mental_state"],
            "tags": ", ".join(r["tags"]),
            "content_text": r["content_text"],
            "source_reference": r["source_reference"],
        })
    return rows


# ---------- phiếu chấm Excel ----------

LIKERT = "1 = Rất kém · 2 = Kém · 3 = Trung bình · 4 = Tốt · 5 = Rất tốt"

SHEETS = {
    "finetune": {
        "title": "1_FineTune",
        "desc": "Dữ liệu fine-tune hội thoại thấu cảm (finetune.jsonl, 5000 mẫu tư vấn tâm lý đã dịch) — dùng cho Module 1.",
        "fields": [("review_id", 9), ("type", 12), ("user", 45), ("assistant", 60)],
        "likert": ["Thấu cảm & công nhận cảm xúc", "Tự nhiên, giọng 'người bạn', không phán xét",
                   "Chất lượng dịch tiếng Việt", "An toàn (không chẩn đoán / kê đơn / gây hại)"],
        "binary": ["Dùng được để fine-tune? (Đạt/Không đạt)"],
    },
    "generative": {
        "title": "2_Generative_QA",
        "desc": "QA tự sinh: context RAG -> câu trả lời chuyển tiếp từ đồng cảm sang gợi ý dinh dưỡng — dùng cho Module 4.",
        "fields": [("review_id", 9), ("mental_state", 12), ("context", 38), ("user", 35), ("assistant", 60)],
        "likert": ["Thấu cảm & chuyển tiếp mượt", "Lời khuyên dinh dưỡng đúng khoa học",
                   "Phù hợp triệu chứng người dùng", "Trung thành với context (không bịa)"],
        "binary": ["An toàn y khoa (Pass/Fail)", "Có cần chuyển tuyến chuyên môn mà chưa nói? (Có/Không)"],
    },
    "nutrition": {
        "title": "3_RAG_DinhDuong",
        "desc": "Corpus dinh dưỡng (Data-to-Text từ VN FCT 2007 / USDA) — bảng nutrition_support cho Module 3.",
        "fields": [("review_id", 9), ("food_item", 16), ("key_nutrients", 22), ("ontology_effect", 22),
                   ("content_text", 55), ("source_reference", 20)],
        "likert": ["Thực phẩm thực sự giàu các chất đã nêu", "Cơ chế tác động lên tâm lý đúng khoa học",
                   "Câu văn rõ ràng, dễ hiểu"],
        "binary": ["Nhãn ontology_effect đúng? (Đúng/Sai)", "Có phát biểu quá mức / gây hiểu nhầm? (Có/Không)"],
    },
    "psychology": {
        "title": "4_RAG_TamLy",
        "desc": "Corpus tâm lý (CBT – CCI, MBSR, WHO PFA), đã dịch & chunk — bảng psychology_docs cho Module 3.",
        "fields": [("review_id", 9), ("target_mental_state", 16), ("tags", 22), ("content_text", 65),
                   ("source_reference", 22)],
        "likert": ["Nội dung đúng chuyên môn tâm lý", "Chất lượng dịch tiếng Việt",
                   "Chunk đủ ngữ cảnh, đọc độc lập được"],
        "binary": ["Nhãn target_mental_state đúng? (Đúng/Sai)", "Tags phù hợp? (Đúng/Sai)"],
    },
}

BINARY_CHOICES = {
    "Đạt/Không đạt": '"Đạt,Không đạt"', "Pass/Fail": '"Pass,Fail"',
    "Đúng/Sai": '"Đúng,Sai"', "Có/Không": '"Có,Không"',
}

HEAD_FILL = PatternFill("solid", fgColor="2F5496")
RATE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_sheet(wb, spec, rows):
    ws = wb.create_sheet(spec["title"])
    ws["A1"] = spec["desc"]
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Thang Likert: {LIKERT}. Ô màu vàng là ô chuyên gia điền."
    ws["A2"].font = Font(italic=True)

    headers = [f for f, _ in spec["fields"]] + spec["likert"] + spec["binary"] + ["Nhận xét / đề xuất sửa"]
    widths = [w for _, w in spec["fields"]] + [14] * len(spec["likert"]) + [16] * len(spec["binary"]) + [40]
    hr = 4
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[hr].height = 60

    nf = len(spec["fields"])
    nl = len(spec["likert"])
    first, last = hr + 1, hr + len(rows)
    for r_i, row in enumerate(rows, first):
        longest = 0
        for c, (f, w) in enumerate(spec["fields"], 1):
            v = row[f]
            cell = ws.cell(row=r_i, column=c, value=v)
            cell.alignment = WRAP
            cell.border = BORDER
            longest = max(longest, len(str(v)) / max(w, 1))
        for c in range(nf + 1, len(headers) + 1):
            cell = ws.cell(row=r_i, column=c)
            cell.fill = RATE_FILL
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
        ws.row_dimensions[r_i].height = min(409, max(45, 15 * (longest * 1.15 + 1)))

    def col(i):
        return ws.cell(row=hr, column=i).column_letter

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        showErrorMessage=True, error="Nhập số nguyên từ 1 đến 5")
    ws.add_data_validation(dv)
    dv.add(f"{col(nf + 1)}{first}:{col(nf + nl)}{last}")
    for k, b in enumerate(spec["binary"]):
        choice = next(v for key, v in BINARY_CHOICES.items() if key in b)
        dvb = DataValidation(type="list", formula1=choice, allow_blank=True)
        ws.add_data_validation(dvb)
        c = col(nf + nl + 1 + k)
        dvb.add(f"{c}{first}:{c}{last}")

    ws.freeze_panes = ws.cell(row=hr + 1, column=2)


def write_guide(wb, reviewer_system_prompt, counts):
    ws = wb.active
    ws.title = "0_HuongDan"
    lines = [
        ("PHIẾU KIỂM CHỨNG DỮ LIỆU VỚI CHUYÊN GIA", Font(bold=True, size=14)),
        ("Đề tài: Chatbot multi-agent kết hợp RAG hỗ trợ sức khỏe tinh thần cho người Việt Nam bằng thực dưỡng", Font(bold=True)),
        ("A Multi-Agent Chatbot with Stateful Hybrid RAG for Mental Health Support among Vietnamese Users through Nutritional Therapy", Font(italic=True)),
        ("", None),
        ("Họ tên chuyên gia:", Font(bold=True)),
        ("Chuyên môn (Tâm lý / Dinh dưỡng / Y khoa):", Font(bold=True)),
        ("Ngày đánh giá:", Font(bold=True)),
        ("", None),
        ("Cách đánh giá", Font(bold=True, size=12)),
        ("1. Mỗi sheet 1–4 chứa 10 mẫu được lấy ngẫu nhiên phân tầng từ một bộ dữ liệu.", None),
        ("2. Chấm các cột màu vàng: Likert 1–5 (chỉ nhập số nguyên) và các cột lựa chọn (chọn từ danh sách thả xuống).", None),
        (f"3. Thang Likert: {LIKERT}.", None),
        ("4. Ghi rõ lỗi cụ thể và đề xuất sửa ở cột 'Nhận xét / đề xuất sửa' (đặc biệt với điểm ≤ 2 hoặc Fail/Sai).", None),
        ("5. Vui lòng chấm độc lập, không trao đổi với chuyên gia khác (để tính Cohen's Kappa).", None),
        ("6. Nếu nằm ngoài chuyên môn của bạn, để trống ô đó và ghi 'ngoài chuyên môn' ở cột nhận xét.", None),
        ("", None),
        ("Tiêu chí an toàn y khoa (tham chiếu)", Font(bold=True, size=12)),
        ("• Không chẩn đoán bệnh, không kê đơn thuốc, không thay thế điều trị y khoa.", None),
        ("• Không phán xét, không đưa lời khuyên sáo rỗng.", None),
        ("• Có dấu hiệu tự hại / khủng hoảng → phải khuyến khích tìm hỗ trợ chuyên môn / đường dây nóng.", None),
        ("• Lời khuyên dinh dưỡng phải có cơ sở, không phóng đại công dụng.", None),
        ("", None),
        ("Nội dung các sheet", Font(bold=True, size=12)),
    ]
    for key, spec in SHEETS.items():
        lines.append((f"• {spec['title']} ({counts[key]} mẫu): {spec['desc']}", None))
    lines += [("", None), ("System prompt dùng trong finetune.jsonl (áp dụng cho mọi mẫu sheet 1):", Font(bold=True)),
              (reviewer_system_prompt, Font(italic=True))]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font
    ws.column_dimensions["A"].width = 130
    for r in (5, 6, 7):
        ws.cell(row=r, column=2).fill = RATE_FILL
        ws.cell(row=r, column=2).border = BORDER
    ws.column_dimensions["B"].width = 40


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--n", type=int, default=10)
    args = ap.parse_args()
    rng = random.Random(args.seed)

    ft_rows, ft_sys = sample_finetune(rng, args.n)
    data = {
        "finetune": ft_rows,
        "generative": sample_generative(rng, args.n),
        "nutrition": sample_nutrition(rng, args.n),
        "psychology": sample_psychology(rng, args.n),
    }

    meta = {"seed": args.seed, "n_per_dataset": args.n,
            "sources": {k: str(v.relative_to(ROOT)) for k, v in SRC.items()},
            "finetune_system_prompt": ft_sys}
    with open(OUT / "samples.json", "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "samples": data}, f, ensure_ascii=False, indent=2)

    wb = Workbook()
    write_guide(wb, ft_sys, {k: len(v) for k, v in data.items()})
    for key, spec in SHEETS.items():
        write_sheet(wb, spec, data[key])
    wb.save(OUT / "expert_review_sheet.xlsx")

    for k, v in data.items():
        print(k, len(v), [r.get("mental_state") or r.get("target_mental_state") or r.get("ontology_effect") or r.get("type") for r in v])


if __name__ == "__main__":
    main()
