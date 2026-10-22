"""Dựng file Word nhật ký các hướng đã thử và các lần chạy đánh giá, kèm số liệu.

    python scripts/build_experiment_log.py      ->  ../NHAT_KY_THU_NGHIEM.docx

Số của các lần chạy hệ thống được tính lại từ file kết quả (eval_result_afterChangePrompt, eval_result_final,
eval_grounding_result_2, eval_routing_baseline_results); số của các thử nghiệm trong thu-nghiem/ lấy từ thu-nghiem/README.md
và BAO_CAO_DANH_GIA.md (mục 5.1). Báo cáo chính thức vẫn là BAO_CAO_KET_QUA_CUOI.docx.
"""
import json
import os
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AIM = os.path.dirname(ROOT)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import build_final_report as R  # noqa: E402  (dùng lại tbl, bullet, note, pc)
from eval_rules import NO_BODY, STATES, routing  # noqa: E402

OUT = os.path.join(AIM, "NHAT_KY_THU_NGHIEM.docx")
RUN1 = os.path.join(AIM, "eval_result_afterChangePrompt")
RUN2 = os.path.join(AIM, "eval_result_final")
G1 = os.path.join(RUN1, "eval_grounding_results", "logs", "eval", "grounding")
G2 = os.path.join(RUN2, "eval_grounding_results", "logs", "eval", "grounding")
G3 = os.path.join(AIM, "eval_grounding_result_2", "logs", "eval", "grounding")
RB = os.path.join(AIM, "eval_routing_baseline_results", "logs", "eval", "grounding_pipeline_qwen_goc")
jl, tbl, bullet, note, pc = R.jl, R.tbl, R.bullet, R.note, R.pc


def analysis(run_logs):
    return json.load(open(os.path.join(run_logs, "eval", "analysis_final.json"), encoding="utf-8"))


def grounding(g):
    a = jl(os.path.join(g, "A_qwen-huhu_scored.jsonl"))
    ng = [r for x in ("B1", "B2", "B3", "B4") for r in jl(os.path.join(g, f"B_{x}_scored.jsonl")) if "final_grounding" in r]
    ro = routing(g)
    flag = sum(v[1] for k, v in ro["b3_groups"].items() if k not in NO_BODY.values())
    n = sum(v[0] for k, v in ro["b3_groups"].items() if k not in NO_BODY.values())
    return {"a_json": pc(sum(r["json_ok"] for r in a), len(a)), "a_grounded": pc(sum(r["grounded"] for r in a), len(a)),
            "sys_refs": pc(sum(r["final_grounding"]["refs_in_context"] for r in ng), len(ng)),
            "sys_grounded": pc(sum(r["final_grounding"]["grounded"] for r in ng), len(ng)),
            "b2": f"{ro['b2_full_ctx']}/{ro['b2_full'] - ro['b2_full_med']}", "b2_full": ro["b2_full"],
            "b3_med": f"{ro['b3_flag_med']}/{ro['b3_flag']}", "b3_flag": f"{flag}/{n}"}


def compare_baseline(g):
    """So từng hội thoại giữa P (thư mục g) và Pipeline + Qwen2.5-7B gốc (RB), như build_final_report."""
    last = lambda r: r["turns"][-1]
    ok = lambda r: all(last(r)["states"].get(k) for k in STATES) and not last(r)["medication"]
    c = {"both": 0, "both_db": 0, "only": [0, 0], "t1s": 0, "t1b": 0, "n": 0}
    for x in ("B2", "B3"):
        pr = {r["id"]: r for r in jl(os.path.join(g, f"B_{x}_runs.jsonl"))}
        for i, b in ((r["id"], r) for r in jl(os.path.join(RB, f"B_{x}_runs.jsonl"))):
            a = pr[i]
            c["t1s"] += a["turns"][0]["states"] != b["turns"][0]["states"]
            c["t1b"] += a["turns"][0]["branch"] != b["turns"][0]["branch"]
            c["n"] += 1
            if x == "B2":
                if ok(a) and ok(b):
                    c["both"] += 1
                    c["both_db"] += last(a)["branch"] in ("nutrition", "nutrition_blocked") and last(b)["branch"] in ("nutrition", "nutrition_blocked")
                else:
                    c["only"][0] += ok(a)
                    c["only"][1] += ok(b)
    return c


def build():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(12)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[s].font.name = "Times New Roman"
        doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)

    doc.add_heading("Nhật ký thử nghiệm và các lần chạy đánh giá", 0)
    doc.add_paragraph("Tài liệu ghi lại các hướng đã thử, các lần chạy đánh giá trên cùng bộ test và các thay đổi cách chấm, kèm số liệu. "
                      "Kết quả chính thức nằm trong BAO_CAO_KET_QUA_CUOI.docx. Ký hiệu: P = hệ thống đề xuất (pipeline + qwen-huhu); 3 baseline "
                      "gồm Qwen2.5-7B gốc và qwen-huhu trả lời trực tiếp, và Pipeline + Qwen2.5-7B gốc. Mọi số chấm câu trả lời là số chấm tự động.")

    # ---------------------------------------------------------------- 1
    doc.add_heading("1. Dòng thời gian", 1)
    tbl(doc, ["Ngày", "Việc", "Kết quả chính", "Trong báo cáo cuối"], [
        ["16–17/10", "Lần chạy đánh giá đầy đủ (sau khi sửa lỗi cài đặt nhóm cờ y tế B3)", "Crisis P 149/164; safety P 278/300", "Không (thay bằng lần 20/10); chấm tay làm trên lần này"],
        ["17/10", "Nhóm chấm tay các ca P trượt", "Safety 22 câu, crisis 15 ca", "Có (phân tích lỗi)"],
        ["18/10", "Thử nghiệm 1: lớp kiểm tra khủng hoảng thứ hai (giám khảo gemma3:4b sau LLM #1)", "Crisis 149 → 161/164 nhưng 157/958 hội thoại thường bị chuyển sang phản hồi khủng hoảng", "Hướng phát triển"],
        ["19–20/10", "Thử nghiệm 2: LLM #1 mới + LLM #3 (gemma3:12b) + P0 làm baseline, notebook TPU và GPU", "Chạy dở; LLM #1 mới phát hiện 141/164; safety kém P0", "Không"],
        ["20/10", "Thử nghiệm 3: sửa định nghĩa cờ y tế B3 trong prompt LLM #1", "Crisis 145/164, safety 281/300; không thay đổi nào có ý nghĩa thống kê", "Không"],
        ["20/10", "Lần chạy cuối, notebook chính (cùng hệ thống với 16–17/10)", "Crisis P 149/164; safety P 286/300", "Có"],
        ["20/10", "Lần chạy cuối, notebook grounding", "Kiểm tra tích hợp 72/72; cờ y tế 168/224", "Không (thay bằng lần 21/10)"],
        ["21/10", "Hội thoại nhiều lượt của baseline Pipeline + Qwen2.5-7B gốc", "64/64; cờ y tế 168/224; 647/651 lượt là câu dự phòng", "Có"],
        ["21/10", "Notebook grounding chạy lại, thêm chấm hỏi lại state bằng giám khảo LLM", "75/75; cờ y tế 167/224; hỏi lại state 88,0%", "Có (trừ tỷ lệ hỏi lại state)"],
        ["22/10", "Chấm hỏi lại state của P trên bộ crisis/safety (câu trả lời 20/10)", "122/180 (67,8%)", "Không (chỉ số bỏ, mục 5)"],
        ["22/10", "Bỏ 2 chỉ số khỏi báo cáo: tỷ lệ hỏi lại state; cột giám khảo \"tác dụng ngoài dữ liệu\" và tỷ lệ \"Grounded\"", "Xem mục 5", "–"],
    ], widths=[1.9, 6.2, 5.6, 3.5])

    # ---------------------------------------------------------------- 2
    doc.add_heading("2. Các lần chạy hệ thống trên cùng bộ test", 1)
    doc.add_paragraph("Ba lần chạy dùng cùng code, prompt, dữ liệu và model. Số khác nhau do model trả lời chạy temperature 0,3 và LLM #1 dao động nhẹ giữa các lần chạy.")
    a1, a2 = analysis(os.path.join(RUN1, "eval_final_results (1)", "logs")), analysis(os.path.join(RUN2, "eval_final_results", "logs"))
    eff = lambda an, s, su: pc(an[s][su]["effective"], an[s][su]["n"])
    rows = []
    for s, name in (("P", "P"), ("B1", "Qwen2.5-7B gốc"), ("B2", "qwen-huhu"), ("B3", "Pipeline + Qwen2.5-7B gốc")):
        rows.append([f"Crisis đạt hiệu quả: {name}", eff(a1, s, "crisis"), eff(a2, s, "crisis")])
    for s, name in (("P", "P"), ("B1", "Qwen2.5-7B gốc"), ("B2", "qwen-huhu"), ("B3", "Pipeline + Qwen2.5-7B gốc")):
        rows.append([f"Safety đạt hiệu quả: {name}", eff(a1, s, "safety"), eff(a2, s, "safety")])
    rows += [["LLM #1 phát hiện khủng hoảng", "148/164", "148/164"],
             ["Truy xuất theo 4 state, đúng chiều vấn đề", "269/295", "269/295"],
             ["Báo khủng hoảng trên 300 câu safety (P)", "3 (SAFE_095, 185, 244)", "2 (SAFE_185, 244)"],
             ["Ca P trượt: crisis / safety", "15 / 22", "15 (trùng) / 14 (10 trùng)"]]
    tbl(doc, ["Notebook chính", "16–17/10", "20/10 (dùng trong báo cáo)"], rows, widths=[8, 4.5, 4.7])
    g = {k: grounding(v) for k, v in (("16/10", G1), ("20/10", G2), ("21/10", G3))}
    tbl(doc, ["Notebook grounding", "16/10", "20/10", "21/10 (dùng trong báo cáo)"], [
        ["Mức model qwen-huhu: JSON đúng", g["16/10"]["a_json"], g["20/10"]["a_json"], g["21/10"]["a_json"]],
        ["Mức model qwen-huhu: grounded (gồm cột giám khảo)", g["16/10"]["a_grounded"], g["20/10"]["a_grounded"], g["21/10"]["a_grounded"]],
        ["Toàn hệ thống: mã nguồn thuộc context", g["16/10"]["sys_refs"], g["20/10"]["sys_refs"], g["21/10"]["sys_refs"]],
        ["Toàn hệ thống: grounded (gồm cột giám khảo)", g["16/10"]["sys_grounded"], g["20/10"]["sys_grounded"], g["21/10"]["sys_grounded"]],
        ["Nhóm dữ liệu B2: hội thoại đủ 4 state", g["16/10"]["b2_full"], g["20/10"]["b2_full"], g["21/10"]["b2_full"]],
        ["Kiểm tra tích hợp: đủ 4 state, không cờ y tế → có món từ PostgreSQL", g["16/10"]["b2"], g["20/10"]["b2"], g["21/10"]["b2"]],
        ["Cờ y tế bật → nhánh thuốc", g["16/10"]["b3_med"], g["20/10"]["b3_med"], g["21/10"]["b3_med"]],
        ["Cờ bật trên các tình trạng thuộc định nghĩa", g["16/10"]["b3_flag"], g["20/10"]["b3_flag"], g["21/10"]["b3_flag"]],
    ], widths=[7.4, 3.2, 3.2, 3.4])
    note(doc, "Lần 16/10 có 1 mẫu mức model sai định dạng (non_vietnamese_script). Lần 21/10 có thêm bước chấm hỏi lại state (mục 5).")

    # ---------------------------------------------------------------- 3
    doc.add_heading("3. Các thử nghiệm không đưa vào hệ thống", 1)
    doc.add_heading("3.1 Lớp kiểm tra khủng hoảng thứ hai (18/10; thu-nghiem/1_llm3-kiem-tra-khung-hoang)", 2)
    doc.add_paragraph("Thêm một lượt gọi giám khảo gemma3:4b ngay sau LLM #1, ở mọi lượt LLM #1 không bật cờ, để hỏi người dùng có ý định tự sát, "
                      "tự hại hoặc đang gặp nguy hiểm cần giúp gấp không; có thì trả phản hồi khủng hoảng soạn sẵn. So với lần chạy 16–17/10.")
    tbl(doc, ["Chỉ số", "Hệ thống (16–17/10)", "Có lớp kiểm tra thứ hai"], [
        ["Crisis đạt hiệu quả P (164 ca)", "149 (90,9%)", "161 (98,2%); McNemar p = 0,0005"],
        ["Báo khủng hoảng trên 300 câu safety (P)", "3", "9"],
        ["Hội thoại bình thường bị chuyển sang phản hồi khủng hoảng (tập test fine-tune)", "–", "157/958 (16,4%); nhóm B1 109/238"],
    ], widths=[8, 4, 5.2])
    bullet(doc, "Bắt thêm 12 ca khủng hoảng LLM #1 bỏ sót, không mất ca nào; nhưng giám khảo bật cờ với cảm xúc tiêu cực mạnh mà không nguy hiểm "
                "(khóc, uất ức, buồn, căng thẳng), nhất là ở lượt 2. Không đưa vào vì báo nhầm quá cao.")
    doc.add_heading("3.2 LLM #1 mới, LLM #3 kiểm tra câu trả lời, chạy TPU/GPU (19–20/10; thu-nghiem/2_llm1-moi-llm3-p0-tpu-gpu)", 2)
    for t in ["LLM #1 gộp tự sát, tự hại, hoảng loạn cấp thành một cờ khủng hoảng; định nghĩa cờ y tế B3 mới.",
              "LLM #3 (gemma3:12b) kiểm tra câu trả lời cuối, thay guardrail tầng 2. P0 = hệ thống chính thức làm baseline; notebook chạy trên Kaggle TPU (vLLM) và GPU (Ollama).",
              "Kết quả sơ bộ trên GPU (chưa phân tích xong do thiếu 4 ca của baseline qwen-huhu): so với P0, crisis 148 vs 149/164 (p = 1,0), safety 276 vs 285/300 "
              "(P0 tốt hơn, McNemar p = 0,049); LLM #1 mới phát hiện khủng hoảng 141/164 (bản chính thức 148). Không dùng."]:
        bullet(doc, t)
    doc.add_heading("3.3 Sửa định nghĩa cờ y tế B3 trong prompt LLM #1 (20/10; thu-nghiem/3_prompt-b3)", 2)
    doc.add_paragraph("Rượu bia, thuốc lá/vape, cà phê, nhịn ăn, thức trắng chỉ là B3 khi đã gây triệu chứng thể chất. Tập dev 16 câu tự viết: 11/16 → 15/16. "
                      "Chạy trên Kaggle GPU với cùng hạ tầng, model và bộ test; so với lần chạy 16–17/10.")
    tbl(doc, ["Chỉ số (P)", "16–17/10", "Prompt B3", "So cặp"], [
        ["Khủng hoảng đạt hiệu quả", "149/164", "145/164", "0 tốt lên, 4 xấu đi; p = 0,125"],
        ["An toàn đạt hiệu quả", "278/300", "281/300", "12 tốt lên, 9 xấu đi; p = 0,66"],
        ["Truy xuất theo 4 state (đúng chiều)", "269/295", "271/295", ""],
        ["Cờ y tế B3 (tình trạng thuộc định nghĩa)", "169/224", "171/224", ""],
    ], widths=[6, 3, 3, 5.2])
    bullet(doc, "Cờ y tế tốt lên ở rượu bia (5 → 13/21), nhịn ăn có triệu chứng (6 → 12/13), thức trắng có triệu chứng (0 → 5/10), cà phê (2 → 4/9); "
                "xấu đi ở kỳ kinh (15 → 7/19), mang thai (21 → 16/21), cai thuốc lá/vape (6 → 2/16), vừa mổ (19 → 18); bật nhầm 3 ca không kể triệu chứng. "
                "Không thay đổi nào có ý nghĩa thống kê; không đưa vào hệ thống.")

    # ---------------------------------------------------------------- 4
    doc.add_heading("4. Hội thoại nhiều lượt của baseline Pipeline + Qwen2.5-7B gốc (21/10)", 1)
    doc.add_paragraph("Cùng pipeline, chỉ đổi model trả lời; nhóm dữ liệu B2 và B3 của tập test fine-tune (480 hội thoại, 651 lượt). Model gốc trả câu "
                      "dự phòng ở 647/651 lượt. So từng hội thoại với P của hai lần chạy grounding.")
    rb = routing(RB)
    c2, c3 = compare_baseline(G2), compare_baseline(G3)
    r2, r3 = routing(G2), routing(G3)
    tbl(doc, ["Chỉ số", "So với P 20/10", "So với P 21/10 (dùng trong báo cáo)"], [
        ["Hội thoại cả hai cùng đủ 4 state, không cờ y tế → cả hai có món từ PostgreSQL", f"{c2['both_db']}/{c2['both']}", f"{c3['both_db']}/{c3['both']}"],
        ["Chỉ P gom đủ 4 state / chỉ baseline gom đủ", f"{c2['only'][0]} / {c2['only'][1]}", f"{c3['only'][0]} / {c3['only'][1]}"],
        ["Kiểm tra tích hợp: P / baseline", f"{r2['b2_full_ctx']}/{r2['b2_full'] - r2['b2_full_med']} / {rb['b2_full_ctx']}/{rb['b2_full'] - rb['b2_full_med']}",
         f"{r3['b2_full_ctx']}/{r3['b2_full'] - r3['b2_full_med']} / {rb['b2_full_ctx']}/{rb['b2_full'] - rb['b2_full_med']}"],
        ["Lượt 1 (đầu vào giống hệt): state khác / nhánh khác", f"{c2['t1s']}/{c2['n']} / {c2['t1b']}/{c2['n']}", f"{c3['t1s']}/{c3['n']} / {c3['t1b']}/{c3['n']}"],
    ], widths=[8.4, 3.8, 5])
    bullet(doc, "Ví dụ đã dùng để giải thích vì sao mẫu số khác nhau: B2_075 (so với P 20/10), B2_079 (so với P 21/10). Ở B2_075 của lần 21/10, P cũng không gom đủ 4 state.")

    # ---------------------------------------------------------------- 5
    doc.add_heading("5. Thay đổi cách chấm và cách trình bày", 1)
    tbl(doc, ["Thay đổi", "Trước", "Sau"], [
        ["Truy xuất", "So đích tác động với nhãn của bộ test, có nhóm digestion", "Chấm theo 4 state hệ thống nhận vào; 47 câu chỉ có nhãn digestion hoặc impulsivity_control không tính"],
        ["Định tuyến nhiều lượt", "So nhánh với nhãn của tập test fine-tune (nhãn gắn theo kịch bản sinh dữ liệu)", "Chỉ chấm khi thông tin đã có trong hội thoại; chỉ số B2 gọi là kiểm tra tích hợp"],
        ["Nhịn ăn, thức trắng", "Tính là cờ y tế B3", "Chỉ tính khi người dùng kể triệu chứng thể chất"],
        ["Câu dự phòng", "Có cột \"đạt cả 5 tiêu chí (tự động)\" tính cả câu dự phòng: Pipeline + Qwen2.5-7B gốc 98,0% safety", "Câu dự phòng tính là không đạt ở mọi chỉ số: 0,3% safety; cột 98,0% bỏ"],
        ["Tên baseline", "B1, B2, B3 (trùng tên nhóm dữ liệu B1–B4)", "Gọi theo model trả lời; mã B1/B2/B3 chỉ còn trong code và tên file"],
        ["Hỏi lại state còn thiếu", "Dò từ khoá: 132/180 (73,3%) trên bộ crisis/safety", "Giám khảo LLM: 122/180 (67,8%) trên cùng 180 câu; hội thoại nhiều lượt 929/1.056 (88,0%). "
         "Sau đó bỏ khỏi báo cáo: chatbot tâm lý không phải lúc nào cũng cần hỏi; ở 53/58 lượt không hỏi của bộ crisis/safety, bot thấu cảm hoặc khuyên gặp bác sĩ mà không đặt câu hỏi"],
        ["Grounding: cột giám khảo", "\"Tác dụng ngoài dữ liệu\" (21/10: qwen-huhu 51/240 đạt) và \"món ngoài dữ liệu\" (143/240), gộp thành \"Grounded\" 47/240",
         "Bỏ: model bổ sung giải thích từ kiến thức phổ thông về dưỡng chất (transfer learning kiến thức đại chúng; câu bổ sung có ở 231/240 câu trả lời, "
         "học từ 596/892 mẫu train nhóm B2), báo cáo xem là điểm mạnh; cột \"món ngoài dữ liệu\" chấm nhầm món có trong context. Giữ các cột chấm bằng luật"],
        ["Bố cục báo cáo Word", "Theo loại kết quả", "3 phần: toàn hệ thống, truy xuất database, model fine-tune; đánh dấu (core) các mục có đáp án từng ca"],
    ], widths=[3.4, 6.4, 7.4])
    doc.add_heading("Cách chấm hỏi lại state bằng LLM: các phiên bản đã thử", 2)
    tbl(doc, ["Phiên bản", "Cách làm", "Kết quả khi thử (câu trả lời của lần chạy 20/10)"], [
        ["1", "Giám khảo đọc cả câu trả lời, trả danh sách mã state (sleep_state…)", "12 ca: chép sai mã (bot hỏi giấc ngủ mà ghi tâm trạng)"],
        ["2", "Giám khảo đọc cả câu trả lời, trả true/false theo tên tiếng Việt", "12 ca: 4 ca thiếu khoá; coi câu khuyên (\"Bạn hỏi bác sĩ…\") là câu hỏi"],
        ["3 (dùng)", "Tách câu hỏi (\"?\" hoặc lời mời kể, chia sẻ), giám khảo chỉ xếp các câu hỏi vào 4 state", "40 ca: nhóm đọc lại, chấm đúng cả 40"],
    ], widths=[2, 7.6, 7.6])

    # ---------------------------------------------------------------- 6
    doc.add_heading("6. Các lựa chọn thiết kế đã thử", 1)
    tbl(doc, ["Lựa chọn", "Đã thử", "Chọn"], [
        ["Giám khảo (guardrail tầng 2 và chấm điểm), 32 ca tự soạn", "gemma3:4b 30/32 đúng, 2 gắn nhầm; aya-expanse:8b 27/32, 5; qwen2.5:7b 26/32, 6", "gemma3:4b (khác gốc với model fine-tune)"],
        ["Khoá JSON của cờ y tế trong LLM #1", "medication_alert: model bỏ qua phần \"cơ thể khác thường\"", "health_alert (parser nhận cả hai); 10 câu tự viết 9/10 đúng"],
        ["Chạy bge-m3 trên Kaggle", "GPU: Ollama phải đẩy một phần qwen-huhu sang CPU, chạy chậm nhiều lần", "CPU"],
    ], widths=[5, 7.2, 5])

    # ---------------------------------------------------------------- 7
    doc.add_heading("7. Vị trí dữ liệu", 1)
    tbl(doc, ["Nội dung", "Thư mục"], [
        ["Lần chạy 16–17/10", "eval_result_afterChangePrompt/"],
        ["Lần chạy 20/10 (notebook chính và grounding)", "eval_result_final/"],
        ["Notebook grounding 21/10", "eval_grounding_result_2/"],
        ["Baseline Pipeline + Qwen2.5-7B gốc, hội thoại nhiều lượt (21/10)", "eval_routing_baseline_results/"],
        ["Thử nghiệm 1–3 (code, gói chạy, kết quả)", "thu-nghiem/ (thu-nghiem/README.md)"],
        ["Chấm tay (lần chạy 16–17/10)", "Hybrid-database-main/logs/eval/chamtay_after_b3fix.md"],
    ], widths=[8, 9.2], bold_first_col=True)
    doc.save(OUT)
    R._fix_zoom(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
